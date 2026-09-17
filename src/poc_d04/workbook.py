"""스프레드시트 활동지(xlsx): 누락 표시표 · 역할표 · PoC 기획안 8칸 · 5항목 평가표 · 시간 기록지 · 참조표."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .catalog import LEXICON, REQUIRED_FIELDS
from .data_gen import load_dataset
from .pipeline import run_doc

HEAD = PatternFill("solid", fgColor="DDEBF7")
HUMAN = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD, WRAP
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.freeze_panes = "A2"
    return ws


def _guide(wb):
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 110
    lines = [
        "D04 · 구매 요청서 누락 확인 PoC — 활동지 (교육용 가상 자료)",
        "",
        "사용 순서 (M03 공통 시연 8–13분 조별 실습)",
        "1. '표본목록'에서 자기 조 요청서 1건을 고른다. 텍스트는 data/requests/ 파일 또는 오프라인 캡처를 본다.",
        "2. ChatGPT에 prompts/00_시스템규칙 + 01_추출을 넣고 JSON을 받는다. (접속 불가 시 examples/D04_오프라인_추출결과 사용)",
        "3. '누락표시표'에 원문 값 / 표준명 후보(AI 제안) / 상태(규칙) / 근거 / 보완 질문을 옮긴다. 값이 없으면 비운다.",
        "4. '사람 확정' 열은 반드시 사람이 채운다. AI 출력을 그대로 붙이지 않는다.",
        "5. '역할표'에 AI 제안 / 규칙·화면 / 사람 확정을 나눈다.",
        "6. 'PoC기획안_8칸'의 입력 자료·사람 판단·성공/중단 기준을 채운다.",
        "",
        "게이트 (기획안 §3-2)",
        "G1 값이 없으면 비우거나 '확인 필요'. 원문 없는 값을 채우지 않는다.",
        "G2 '다음 주' 같은 모호 표현을 날짜로 단정하지 않는다.",
        "G3 별칭이 비슷하다는 이유만으로 다른 규격을 합치지 않는다.",
        "G4 원문에 없는 규격을 채우지 않는다.",
        "G5 최종 확정은 사람이 한다.",
        "",
        "판단 기준: '문서를 만들었는가'가 아니라 '추가 문의 항목(누락·모호)을 찾았는가'.",
        "모든 품목 ID·규격·현장명은 교육용 가상이며 실제 건설 자재 규격을 뜻하지 않는다. 표본 수·목표 %는 교육용 가정이다.",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=13)
    for r in (3, 11):
        ws.cell(row=r, column=1).font = BOLD


def _table_sheet(wb, dataset, sample_ids):
    headers = ["문서ID", "항목", "원문 값", "원문 발췌 위치", "표준명 후보 (AI 제안)", "상태 (규칙)", "확인 필요",
               "근거", "보완 질문 초안 (AI 제안)", "사람 확정", "확정자", "비고"]
    ws = _sheet(wb, "누락표시표", headers, [10, 10, 22, 12, 30, 12, 10, 34, 46, 22, 10, 18])
    dv = DataValidation(type="list", formula1='"충족,누락,모호,불일치"', allow_blank=True)
    ws.add_data_validation(dv)
    row = 2
    for doc_id in sample_ids:
        text = dataset["문서"][doc_id]["원문"]
        res = run_doc(text, doc_id, "rule")
        for r in res["판정"]["항목"]:
            pos = f"{r['원문시작']}–{r['원문끝']}" if r["원문시작"] is not None else ""
            ws.append([doc_id, r["항목명"], r["원문값"] if r["원문값"] is not None else "", pos,
                       r["표준명후보"] or "", r["상태"], f'=IF(F{row}="충족","","확인 필요")', r["근거"],
                       r["보완질문"] or "", "", "", ""])
            dv.add(f"F{row}")
            ws.cell(row=row, column=10).fill = HUMAN
            for c in range(1, 13):
                ws.cell(row=row, column=c).alignment = WRAP
            row += 1
        row += 1
    # 빈 양식 6항목 × 2건
    for _ in range(2):
        for name in ["품목 ID", "수량", "단위", "희망일", "배송 위치", "규격"]:
            ws.append(["(조별 입력)", name, "", "", "", "", f'=IF(F{row}="충족","","확인 필요")', "", "", "", "", ""])
            dv.add(f"F{row}")
            ws.cell(row=row, column=10).fill = HUMAN
            row += 1
        row += 1
    return ws


def _roles_sheet(wb):
    ws = _sheet(wb, "역할표", ["업무 단계", "AI 제안 (문서 읽기·분류·초안)", "규칙·화면 (계산식·필수항목표·공동 목록)", "사람 확정 (판단·승인)", "우리 조 메모"],
                [24, 40, 40, 40, 30])
    rows = [
        ["요청서에서 항목 읽기", "원문 값·발췌 추출, 표준명 후보 제안", "6개 필수 항목 스키마", "읽기 어려운 글자, 원문과 다른 추출 판단"],
        ["자재 표현 정리", "별칭 → 표준명 후보 나열", "승인된 용어 사전", "동일·대체 품목, 현장 적합성 확정"],
        ["누락·모호 찾기", "보완 질문 초안", "필수 항목 규칙표(충족/누락/모호/불일치)", "필수 정의·예외 승인"],
        ["수량·단위 확인", "—", "표준 단위표, 동의어(EA=개) 환산", "포장 단위·가격 조건 유효성"],
        ["희망일 확인", "모호 표현을 '확인 필요'로 표시(날짜 변환 금지)", "기준일 이전 날짜 → 불일치", "실제 납기 협의·확정"],
        ["배송 위치 확인", "현장명·세부 위치 분리", "세부 위치 토큰(동·층·창고·게이트)", "실입고 위치 확정"],
    ]
    for r in rows:
        ws.append(r + [""])
    for _ in range(3):
        ws.append(["", "", "", "", ""])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        row[3].fill = HUMAN


def _plan_sheet(wb):
    ws = _sheet(wb, "PoC기획안_8칸", ["칸", "작성 안내", "예시 (후보 A · 교육용)", "우리 조 작성"], [18, 40, 60, 60])
    rows = [
        ["1 문제 정의", "누가·무엇을·얼마나 자주 반복하는가", "현장지원팀이 불완전한 구매 요청 때문에 건당 평균 2회 추가 문의(교육용 가정). 구매팀은 단위·배송 조건을 다시 확인", ""],
        ["2 범위", "한 자재군 · 한 업무 단계", "배관 자재군 · 구매 요청서 접수 단계. 견적 비교·발주는 제외", ""],
        ["3 입력 자료", "어떤 문서, 몇 건, 누가 제공", "가상 구매 요청서 30건(읽기 쉬운 + 저품질 스캔), 필수 항목표, 용어 사전. 정답표는 현장지원 담당자가 검토", ""],
        ["4 기대 결과", "산출물 형태", "누락 표시표(원문 값/표준명 후보/확인 필요) + 보완 질문 초안 1건당 1개", ""],
        ["5 AI/규칙/사람", "역할 구분", "AI: 추출·후보·질문 초안 / 규칙: 필수 항목 판정·게이트 / 사람: 확정·예외 승인", ""],
        ["6 기준선", "현재 시간·오류를 어떻게 측정", "시간 기록지로 수동 처리 5건 측정(열기·입력·검토·수정 포함). 현재 추가 문의 횟수 기록", ""],
        ["7 성공·중단 기준", "수치 목표와 중단 조건", "탐지율 ≥95%, 오탐 ≤10%, 중요 오류 0, 추적 100%, 시간 중앙값 30% 단축(교육용 가정). 검토 비용으로 총시간이 늘면 중단", ""],
        ["8 실행 계획", "담당·일정·정답 검토자·갱신 책임", "2주 시험: 1주 자료 준비·정답표, 1주 시험·측정. 정답 검토자 1명, 용어 사전 갱신 담당 1명(교육용 가정)", ""],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        row[3].fill = HUMAN
        ws.row_dimensions[row[0].row].height = 60


def _score_sheet(wb):
    ws = _sheet(wb, "5항목평가표", ["평가 항목", "0점", "1점", "2점", "점수(0–2)", "메모"], [22, 30, 30, 30, 10, 30])
    rows = [
        ["문제 정의", "기능 이름만 있음", "문제는 있으나 빈도·담당 없음", "누가·무엇을·얼마나 자주가 있음"],
        ["입력 자료", "없음", "종류만 있음", "종류·건수·제공자·정답 검토자 있음"],
        ["사람 판단", "AI가 다 함", "일부 구분", "AI/규칙/사람 역할이 항목별로 구분"],
        ["검증 기준", "없음", "성공 기준만", "성공·중단 기준과 기준선 측정 방법"],
        ["실행 가능성", "담당·일정 없음", "일정만", "담당·일정·갱신 책임 있음"],
    ]
    for r in rows:
        ws.append(r + ["", ""])
    ws.append(["합계", "", "", "", "=SUM(E2:E6)", "목표 8점. 입력 자료·사람 판단·검증 기준 중 0점이면 보완"])
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, 7):
        dv.add(f"E{r}")
        ws.cell(row=r, column=5).fill = HUMAN
    ws.cell(row=7, column=1).font = BOLD
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP


def _time_sheet(wb):
    ws = _sheet(wb, "시간기록지", ["문서ID", "방식(수동/AI보조)", "열기·읽기(분)", "입력·추출(분)", "검토(분)", "수정·재작업(분)", "합계(분)", "누락 찾은 수", "추가 문의 필요 수", "기록자"],
                [10, 14, 12, 12, 10, 14, 10, 12, 14, 10])
    for i in range(2, 12):
        ws.append(["", "", "", "", "", "", f"=SUM(C{i}:F{i})", "", "", ""])
    ws.append([])
    ws.append(["중앙값(수동)", "", "", "", "", "", '=IFERROR(MEDIAN(IF(B2:B11="수동",G2:G11)),"")'])
    ws.append(["중앙값(AI보조)", "", "", "", "", "", '=IFERROR(MEDIAN(IF(B2:B11="AI보조",G2:G11)),"")'])
    ws.append(["단축률", "", "", "", "", "", '=IFERROR(1-G14/G13,"")'])
    ws.append(["", "시간에는 사람의 검토·수정 시간을 반드시 포함한다(기획안 다섯 문장 ⑤). 배열 수식은 Excel 365/Google Sheets에서 동작."])


def _metrics_sheet(wb, out_dir: Path):
    ws = _sheet(wb, "검증지표표", ["지표", "측정 요지", "교육용 목표", "규칙 백엔드 실측(가상 표본)", "우리 조 PoC 목표"], [22, 50, 18, 22, 20])
    m = None
    p = out_dir / "metrics_rule.json"
    if p.exists():
        m = json.loads(p.read_text(encoding="utf-8"))["전체"]

    def pct(k):
        return "" if not m or m.get(k) is None else f"{m[k]*100:.1f}%"

    rows = [
        ["건당 실제 작업시간", "열기·입력·검토·수정·재작업 포함(대기 중 실작업만)", "중앙값 30% 이상 단축", "시간기록지로 측정", ""],
        ["누락 탐지율", "찾은 필수 누락 ÷ 정답표 필수 누락(항목 수)", "95% 이상", pct("누락 탐지율"), ""],
        ["잘못된 누락 경고 비율", "오탐 ÷ 전체 누락 경고", "10% 이하", pct("잘못된 누락 경고 비율"), ""],
        ["중요 오류", "자재·수량·단위·희망일을 잘못 읽고 확정 가능으로 표시", "시험 표본 0건(소표본≠운영 무오류)", "" if not m else f"{m['중요 오류']}건", ""],
        ["원문 추적 가능 비율", "원문 위치 확인 가능한 추출 값 비율", "100%", pct("원문 추적 가능 비율"), ""],
    ]
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["수치 목표는 교육용 가정이다. 실제 기획에서는 기준선 측정 후 조정한다."])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP


def _ref_sheets(wb, data_dir: Path, dataset):
    ws = _sheet(wb, "필수항목표", list(REQUIRED_FIELDS[0].keys()), [8, 16, 10, 10, 50, 50, 8])
    for r in REQUIRED_FIELDS:
        ws.append(list(r.values()))
    ws2 = _sheet(wb, "용어사전", ["품목ID", "표준명", "규격", "표준단위", "별칭", "주의"], [8, 16, 8, 8, 44, 50])
    for it in LEXICON:
        ws2.append([it.item_id, it.name, it.spec, it.unit, "; ".join(it.aliases), it.note])
    ws3 = _sheet(wb, "표본목록", ["문서ID", "유형", "버전", "파일", "원문"], [10, 12, 8, 44, 80])
    for doc_id, d in dataset["문서"].items():
        ws3.append([doc_id, d["유형"], d["버전"], d["파일"], d["원문"]])
    for w in (ws, ws2, ws3):
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP


def build_workbook(path: Path, data_dir: Path, out_dir: Path, sample_ids=("D04-13",)) -> Path:
    dataset = load_dataset(data_dir)
    wb = Workbook()
    _guide(wb)
    _table_sheet(wb, dataset, sample_ids)
    _roles_sheet(wb)
    _plan_sheet(wb)
    _score_sheet(wb)
    _time_sheet(wb)
    _metrics_sheet(wb, out_dir)
    _ref_sheets(wb, data_dir, dataset)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
