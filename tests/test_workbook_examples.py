from openpyxl import load_workbook

from poc_d04.data_gen import build_dataset
from poc_d04.examples import write_all
from poc_d04.pipeline import run_all
from poc_d04.workbook import build_workbook


def test_workbook_sheets_and_human_column(tmp_path):
    data = tmp_path / "data"
    build_dataset(data)
    run_all(data, tmp_path / "out", "rule")
    p = build_workbook(tmp_path / "out" / "wb.xlsx", data, tmp_path / "out")
    wb = load_workbook(p)
    assert {"안내", "누락표시표", "역할표", "PoC기획안_8칸", "5항목평가표", "시간기록지", "검증지표표",
            "필수항목표", "용어사전", "표본목록"} <= set(wb.sheetnames)
    ws = wb["누락표시표"]
    headers = [c.value for c in ws[1]]
    assert "사람 확정" in headers and "원문 값" in headers
    col = headers.index("사람 확정") + 1
    for row in ws.iter_rows(min_row=2):
        assert row[col - 1].value in (None, ""), "사람 확정 열은 비어 있어야 한다"
    demo_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "D04-13"]
    assert len(demo_rows) == 6
    metrics = wb["검증지표표"]
    assert any("100.0%" in str(c.value) for row in metrics.iter_rows() for c in row)


def test_examples_bad_draft_hits_four_gates(tmp_path):
    data = tmp_path / "data"
    build_dataset(data)
    paths = write_all(data, tmp_path / "ex")
    md = (tmp_path / "ex" / "D04_잘못된초안_예시.md").read_text(encoding="utf-8")
    for code in ("G1", "G2", "G4", "G5"):
        assert f"| {code} |" in md
    cap = tmp_path / "ex" / "D04_오프라인_추출결과"
    assert (cap / "README.md").exists() and len(list(cap.glob("*조_*.md"))) == 6
    assert any(p.name == "시연_D04-13.md" for p in paths)
